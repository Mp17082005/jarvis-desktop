"""Real-desktop UFO verification through Jarvis's own runner (companion.ufo.UFORunner).

Each scenario is judged by independent evidence, never by UFO's self-report alone.
Usage: python -E tests/live_ufo_check.py <scenario>   (excel | youtube | notepad)
"""
import json
import os
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from companion.ufo import UFORunner  # noqa: E402

UFO_HOME = os.environ["JARVIS_UFO_HOME"]
WS = ROOT / "workspace" / "desktop"
REPORT = ROOT / "data" / "ufo-live-report.json"


def run(request, timeout=600):
    runner = UFORunner(home=UFO_HOME, enabled=True, timeout=timeout)
    start = time.time()
    result = runner.run({"request": request})
    result["duration_s"] = round(time.time() - start, 1)
    return result


def excel():
    path = WS / "ufo-sales.xlsx"
    from openpyxl import load_workbook
    before = path.stat().st_mtime
    request = (f"Open the spreadsheet file {path} in the spreadsheet application (LibreOffice Calc). "
               "In sheet Sales, type the header Total in cell D1, and in cells D2, D3 and D4 enter the value of Qty multiplied by Price for that row "
               "(you may type formulas like =B2*C2). Save the file keeping the current xlsx format (choose Keep Current Format if asked), then close the application.")
    result = run(request)
    wb = load_workbook(path, data_only=False)
    ws = wb["Sales"]
    cells = {c: ws[c].value for c in ("D1", "D2", "D3", "D4")}
    # Formulas or numbers both acceptable; evaluate expected values independently.
    expected = {"D2": 20, "D3": 45, "D4": 80}
    def value(v, row):
        if isinstance(v, (int, float)):
            return v
        if isinstance(v, str) and v.replace(" ", "").upper() == f"=B{row}*C{row}":
            return ws[f"B{row}"].value * ws[f"C{row}"].value
        return None
    ok = (path.stat().st_mtime > before and str(cells["D1"]).strip().lower() == "total"
          and all(value(cells[c], int(c[1:])) == v for c, v in expected.items()))
    return {"ufo": result, "evidence": {"cells": cells, "file_modified": path.stat().st_mtime > before}, "verified": ok}


def youtube():
    query = "lofi hip hop"
    request = (f"Open Google Chrome, go to youtube.com, search for {query}, click the first video result (not a Short) "
               "and make sure it starts playing. If an ad plays and a Skip button appears, click Skip; if the Skip "
               "button is not clickable yet, wait a few seconds and look again. Leave the video playing.")
    result = run(request)
    time.sleep(5)
    # Independent evidence: a Chrome window whose title is a YouTube watch page.
    ps = ("Get-Process chrome -ErrorAction SilentlyContinue | Where-Object { $_.MainWindowTitle -like '*YouTube*' } "
          "| Select-Object -ExpandProperty MainWindowTitle")
    titles = subprocess.run(["powershell.exe", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=30).stdout.split("\n")
    titles = [t.strip() for t in titles if t.strip()]
    verified = any(t.endswith("- YouTube - Google Chrome") and not t.startswith("YouTube -") for t in titles)
    return {"ufo": result, "evidence": {"youtube_windows": titles}, "verified": verified}


def notepad():
    out = WS / f"ufo-note-{int(time.time())}.txt"
    text = f"Jarvis UFO check {int(time.time()) % 10000}"
    request = f"Open Notepad, type exactly: {text} , then Save As with the full path {out} in the filename box, then close Notepad."
    result = run(request)
    data = out.read_text(encoding="utf-8", errors="replace").strip() if out.exists() else None
    return {"ufo": result, "evidence": {"file": str(out), "content": data}, "verified": data == text}


if __name__ == "__main__":
    name = sys.argv[1]
    outcome = {"excel": excel, "youtube": youtube, "notepad": notepad}[name]()
    report = json.loads(REPORT.read_text()) if REPORT.exists() else {}
    report[name] = {**outcome, "at": time.strftime("%Y-%m-%d %H:%M:%S")}
    REPORT.write_text(json.dumps(report, indent=2, default=str))
    print(json.dumps(outcome, indent=2, default=str))
    sys.exit(0 if outcome["verified"] else 1)
