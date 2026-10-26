"""Local-only allowlisted desktop actions. Caller owns authentication/approval."""
from pathlib import Path
import importlib.util
import os
import subprocess


APP_EXES = {'chrome': 'chrome.exe', 'edge': 'msedge.exe', 'excel': 'EXCEL.EXE', 'notepad': 'notepad.exe', 'calculator': 'calc.exe', 'explorer': 'explorer.exe'}


def discover_apps():
    """Registry and fixed installation locations only, never PATH/model paths."""
    found = dict.fromkeys(APP_EXES)
    if os.name != 'nt':
        return found
    import winreg
    relatives = {'chrome': ['Google/Chrome/Application/chrome.exe'], 'edge': ['Microsoft/Edge/Application/msedge.exe'], 'excel': ['Microsoft Office/root/Office16/EXCEL.EXE', 'Microsoft Office/Office16/EXCEL.EXE']}
    for app, exe in APP_EXES.items():
        candidates = []
        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            for view in (winreg.KEY_WOW64_64KEY, winreg.KEY_WOW64_32KEY):
                try:
                    with winreg.OpenKey(hive, 'SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\App Paths\\' + exe, 0, winreg.KEY_READ | view) as key:
                        value, kind = winreg.QueryValueEx(key, '')
                        if kind in (winreg.REG_SZ, winreg.REG_EXPAND_SZ):
                            candidates.append(Path(os.path.expandvars(value.strip('"'))))
                except OSError:
                    pass
        for root in ('ProgramFiles', 'ProgramFiles(x86)', 'LOCALAPPDATA'):
            if os.environ.get(root):
                candidates.extend(Path(os.environ[root]) / p for p in relatives.get(app, []))
        windir = Path(os.environ.get('SystemRoot', 'C:/Windows'))
        if app in ('notepad', 'calculator'):
            candidates.append(windir / 'System32' / exe)
        if app == 'explorer':
            candidates.append(windir / exe)
        for candidate in candidates:
            if candidate.is_absolute() and candidate.name.casefold() == exe.casefold() and candidate.is_file() and not candidate.is_symlink():
                found[app] = str(candidate.resolve())
                break
    return found


def _schema(name, properties, required, description):
    return {'type': 'function', 'function': {'name': name, 'description': description,
        'parameters': {'type': 'object', 'properties': properties, 'required': required, 'additionalProperties': False}}}


ACTIONS = [_schema('open_app', {'app': {'type': 'string', 'enum': ['chrome', 'edge', 'excel', 'notepad', 'calculator', 'explorer']}}, ['app'], 'Open an installed allowlisted application on this machine.')]
_TEXT = {'type': 'string'}
_ROWS = {'type': 'array', 'maxItems': 1000, 'items': {'type': 'array', 'maxItems': 1000, 'items': {'type': ['string', 'number', 'boolean', 'null']}}}
for _name in ('create_workbook', 'read_range', 'write_range'):
    _properties = {'path': dict(_TEXT), 'sheet': dict(_TEXT)}
    _properties['range' if _name == 'read_range' else 'rows'] = dict(_TEXT) if _name == 'read_range' else _ROWS
    if _name == 'write_range':
        _properties['start_cell'] = dict(_TEXT)
    ACTIONS.append(_schema(_name, _properties, list(_properties), 'Workbook operation confined to companion workspace; at most 1000 cells/10000 characters. Formulas unsupported.'))
ACTIONS.append(_schema('play_youtube', {'query': {'type': 'string', 'minLength': 1, 'maxLength': 200}}, ['query'], 'Search YouTube and play the first visible non-Short video in the companion browser.'))
ACTIONS.append(_schema('set_volume', {'volume': {'type': 'integer', 'minimum': 0, 'maximum': 100}}, ['volume'], 'Set companion YouTube player volume, not Windows volume.'))
for _name in ('pause', 'resume', 'skip_ad', 'stop', 'capabilities'):
    ACTIONS.append(_schema(_name, {}, [], 'Control companion player or discover this installed machine capabilities.'))
ACTIONS.append(_schema('ufo_task', {'request': {'type': 'string', 'minLength': 1, 'maxLength': 2000}}, ['request'], 'Run one bounded UFO desktop task on this installed machine. Exact approval required; screenshots and request go to the configured UFO provider. Completion is not guaranteed; poll the durable job and never automatically retry.'))
ACTION_RISKS = {s['function']['name']: s['function']['name'] in {'create_workbook', 'write_range', 'play_youtube', 'ufo_task'} for s in ACTIONS}


def action_risk(name):
    return ACTION_RISKS.get(name, True)


class ActionEngine:
    def __init__(self, workspace):
        self.workspace = Path(workspace).absolute()
        from .ufo import UFORunner
        self._ufo = UFORunner()

    def capabilities(self):
        return {'scope': 'installed_machine', 'apps': {name: {'available': path is not None, 'path': path} for name, path in discover_apps().items()}, 'workbooks': bool(importlib.util.find_spec('openpyxl')), 'youtube': bool(importlib.util.find_spec('playwright')), 'ufo': self._ufo.capabilities()}

    def _safe_path(self, value):
        import re
        if not isinstance(value, str) or len(value) > 500 or not value:
            raise ValueError('Invalid path')
        raw = Path(value)
        if '..' in raw.parts or any(':' in part for part in raw.parts[1:] if raw.is_absolute()) or (not raw.is_absolute() and ':' in value):
            raise ValueError('Path traversal or stream')
        path = raw if raw.is_absolute() else self.workspace / raw
        root = self.workspace.resolve()
        if not path.resolve().is_relative_to(root):
            raise ValueError('Outside workspace')
        for item in (path, *path.parents):
            if item.is_symlink() or (hasattr(item, 'is_junction') and item.is_junction()):
                raise ValueError('Links are unsupported')
            if item.exists() and item.is_file() and item.stat().st_nlink != 1:
                raise ValueError('Hard links are unsupported')
        for part in path.relative_to(self.workspace).parts:
            if part.endswith((' ', '.')) or re.match(r'^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)', part, re.I):
                raise ValueError('Reserved Windows name')
        if path.suffix.lower() != '.xlsx' or not path.parent.is_dir():
            raise ValueError('Only xlsx in an existing workspace directory')
        return path

    @staticmethod
    def _bounds(region):
        import re
        from openpyxl.utils.cell import range_boundaries
        if not isinstance(region, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}(?::[A-Z]{1,3}[1-9][0-9]{0,6})?', region):
            raise ValueError('Invalid range')
        c1, r1, c2, r2 = range_boundaries(region)
        if not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576) or (c2-c1+1)*(r2-r1+1) > 1000:
            raise ValueError('Range exceeds bounds')
        return c1, r1, c2, r2

    @staticmethod
    def _rows(rows):
        import math
        if not isinstance(rows, list) or not rows or any(not isinstance(row, list) or not row for row in rows):
            raise ValueError('Rows must be nonempty lists')
        if len({len(row) for row in rows}) != 1 or sum(map(len, rows)) > 1000:
            raise ValueError('At most 1000 rectangular cells')
        chars = 0
        for row in rows:
            for value in row:
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise ValueError('Unsupported value')
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError('Nonfinite value')
                if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
                    raise ValueError('Formula-like strings are unsupported')
                chars += len(str(value))
        if chars > 10000:
            raise ValueError('At most 10000 characters')
        return rows

    def _workbook(self, name, args):
        import tempfile
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
        required = {'path', 'sheet', 'range' if name == 'read_range' else 'rows'}
        if name == 'write_range':
            required.add('start_cell')
        if not isinstance(args, dict) or set(args) != required:
            raise ValueError('Invalid workbook arguments')
        path = self._safe_path(args['path'])
        sheet = args['sheet']
        if not isinstance(sheet, str) or not sheet or len(sheet) > 31:
            raise ValueError('Invalid sheet')
        if name == 'read_range':
            self._bounds(args['range'])
        else:
            self._rows(args['rows'])
            if name == 'write_range':
                from openpyxl.utils.cell import get_column_letter
                c, r, _, _ = self._bounds(args['start_cell'])
                if ':' in args['start_cell']:
                    raise ValueError('Start must be one cell')
                self._bounds(f'{args["start_cell"]}:{get_column_letter(c+len(args["rows"][0])-1)}{r+len(args["rows"])-1}')
        if path.exists() and path.stat().st_size > 10 * 1024 * 1024:
            raise ValueError('Workbook exceeds 10MB')
        if name != 'create_workbook':
            import zipfile
            with zipfile.ZipFile(path) as archive:
                if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Workbook expands beyond bounds')
        if name == 'create_workbook':
            book = Workbook()
            book.active.title = sheet
        else:
            book = load_workbook(path, keep_links=False)
        try:
            ws = book[sheet]
            if name == 'read_range':
                c1, r1, c2, r2 = range_boundaries(args['range'])
                rows = [[cell.value.isoformat() if hasattr(cell.value, 'isoformat') else cell.value for cell in row] for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)]
                if sum(len(str(v)) for row in rows for v in row) > 10000:
                    raise ValueError('Result exceeds 10000 characters')
                return {'ok': True, 'rows': rows, 'path': str(path), 'sheet': sheet}
            r, c = (1, 1) if name == 'create_workbook' else coordinate_to_tuple(args['start_cell'])
            for dr, row in enumerate(args['rows']):
                for dc, value in enumerate(row):
                    ws.cell(r + dr, c + dc).value = value
            fd, temporary = tempfile.mkstemp(suffix='.tmp', dir=path.parent)
            os.close(fd)
            try:
                book.save(temporary)
                if name == 'create_workbook':
                    # Atomic, exclusive publication: unlike replace, never clobbers.
                    os.link(temporary, path)
                else:
                    os.replace(temporary, path)
            finally:
                Path(temporary).unlink(missing_ok=True)
            return {'ok': True, 'path': str(path), 'sheet': sheet, 'status': 'saved'}
        finally:
            book.close()

    def close(self):
        self._ufo.cancel()  # stop any active UFO process tree before engine teardown
        browser = getattr(self, '_youtube', None)
        if browser is not None:
            del self._youtube
            browser.close()

    def execute_cancellable(self, name, args, cancel_check):
        if name == 'ufo_task':
            return self._ufo.run(args, cancel_check=cancel_check)
        return self.execute(name, args)

    def execute(self, name, args):
        if name == 'ufo_task':
            return self._ufo.run(args)
        if name in {'play_youtube', 'pause', 'resume', 'set_volume', 'skip_ad', 'stop'}:
            expected = {'query'} if name == 'play_youtube' else {'volume'} if name == 'set_volume' else set()
            if not isinstance(args, dict) or set(args) != expected:
                return {'ok': False, 'error': 'invalid_arguments'}
            if name == 'play_youtube' and (not isinstance(args['query'], str) or not args['query'].strip() or len(args['query']) > 200):
                return {'ok': False, 'error': 'invalid_arguments'}
            if name == 'set_volume' and (type(args['volume']) is not int or not 0 <= args['volume'] <= 100):
                return {'ok': False, 'error': 'invalid_arguments'}
            try:
                if not hasattr(self, '_youtube'):
                    from .youtube import YouTubeController
                    self._youtube = YouTubeController(self.workspace / 'data' / 'desktop-browser')
                return self._youtube.execute(name, args)
            except ImportError:
                return {'ok': False, 'error': 'unavailable', 'message': 'Playwright is not installed.'}
            except Exception:
                return {'ok': False, 'error': 'browser_failed', 'message': 'Browser action failed or timed out; do not automatically retry.'}
        if name in {'create_workbook', 'read_range', 'write_range'}:
            try:
                return self._workbook(name, args)
            except ImportError:
                return {'ok': False, 'error': 'unavailable', 'message': 'openpyxl is not installed.'}
            except Exception:
                return {'ok': False, 'error': 'workbook_rejected', 'message': 'Workbook operation failed or arguments were rejected.'}
        if name == 'capabilities' and args == {}:
            return {'ok': True, **self.capabilities()}
        if name == 'open_app' and isinstance(args, dict) and set(args) == {'app'} and isinstance(args['app'], str) and args['app'] in APP_EXES:
            path = discover_apps().get(args['app'])
            if not path:
                return {'ok': False, 'error': 'unavailable', 'message': 'Application is not installed on this machine.'}
            try:
                process = subprocess.Popen([path], shell=False)
                return {'ok': True, 'app': args['app'], 'pid': process.pid, 'status': 'launch_requested', 'message': 'Launch requested; window readiness is not verified.'}
            except OSError:
                return {'ok': False, 'error': 'launch_failed', 'message': 'The operating system rejected the launch.'}
        return {'ok': False, 'error': 'unsupported_action', 'message': 'Unsupported action or arguments.'}
